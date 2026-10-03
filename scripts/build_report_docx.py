from pathlib import Path
import re, sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
report=ROOT/'REPORT.md'
out=ROOT/'REPORT.docx'
svg=ROOT/'docs'/'architecture.svg'
png=ROOT/'docs'/'architecture.png'
try:
 import cairosvg
 cairosvg.svg2png(url=str(svg),write_to=str(png),output_width=1400,output_height=720)
except Exception:
 png=None


def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def set_cell_text(cell,text,bold=False,color=None,size=8.5):
    cell.text=''
    p=cell.paragraphs[0]
    r=p.add_run(text)
    r.bold=bold; r.font.size=Pt(size)
    if color: r.font.color.rgb=RGBColor(*color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_table(doc, lines):
    rows=[]
    for ln in lines:
        parts=[p.strip() for p in ln.strip().strip('|').split('|')]
        rows.append(parts)
    if len(rows)>=2 and all(set(x.replace('-','').replace(':','').strip())==set() for x in rows[1]):
        rows=rows[2:]
    elif len(rows)>=2 and all(re.fullmatch(r':?-{3,}:?', x.strip()) for x in rows[1]):
        rows=rows[:1]+rows[2:]
    if not rows: return
    t=doc.add_table(rows=1, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
    for j,h in enumerate(rows[0]):
        set_cell_text(t.rows[0].cells[j],h,True,(255,255,255),8.5); shade(t.rows[0].cells[j],'9E151D')
    for row in rows[1:]:
        cells=t.add_row().cells
        for j in range(len(cells)):
            txt=row[j] if j<len(row) else ''
            set_cell_text(cells[j],txt,False,(35,35,35),8.2); shade(cells[j],'F3EEE8' if len(t.rows)%2==0 else 'FFFFFF')
    doc.add_paragraph()

def add_code(doc, text):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.space_before=Pt(2)
    r=p.add_run(text); r.font.name='Consolas'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(240,240,240)
    pPr=p._p.get_or_add_pPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'141419'); pPr.append(shd)
    return p

doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.6); styles['Normal'].paragraph_format.space_after=Pt(5)
for name,size,color in [('Title',30,'EF242B'),('Heading 1',18,'A90F17'),('Heading 2',13,'222222'),('Heading 3',10.5,'A90F17')]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(color)


def add_inline(p, text):
    parts=re.split(r'(\*\*[^*]+\*\*|`[^`]+`)', text)
    for part in parts:
        if not part: continue
        if part.startswith('**') and part.endswith('**'):
            r=p.add_run(part[2:-2]); r.bold=True
        elif part.startswith('`') and part.endswith('`'):
            r=p.add_run(part[1:-1]); r.font.name='Consolas'; r.font.size=Pt(8.8)
        else:
            p.add_run(part)

# Header/footer
header=sec.header.paragraphs[0]; header.text='NOVAMART GUARDIAN  /  MANTRA YUDHA'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(169,15,23)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; rr=footer.add_run('Competition build report  |  v1.0.0'); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor(120,120,120)

# Title page
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(100)
r=p.add_run('NOVAMART GUARDIAN'); r.font.size=Pt(34); r.font.bold=True; r.font.color.rgb=RGBColor(239,36,43)
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Mantra Yudha - AI Customer Support Agent from Hell'); r.font.size=Pt(16); r.font.bold=True
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Complete Project Report, Architecture, Implementation and Evaluation'); r.font.size=Pt(11); r.font.color.rgb=RGBColor(110,110,110)
if png and png.exists():
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; doc.add_picture(str(png),width=Inches(6.7))
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run('Understand -> Verify -> Reason -> Decide -> Act / Ask / Escalate'); r.bold=True; r.font.color.rgb=RGBColor(169,15,23)
doc.add_page_break()

lines=report.read_text(encoding='utf-8').splitlines()
i=0; table_buf=[]; code=False; code_buf=[]
while i<len(lines):
    ln=lines[i]
    if ln.startswith('```'):
        if not code:
            code=True; code_buf=[]
        else:
            code=False; add_code(doc,'\n'.join(code_buf))
        i+=1; continue
    if code:
        code_buf.append(ln); i+=1; continue
    if ln.startswith('|'):
        table_buf.append(ln); i+=1
        while i<len(lines) and lines[i].startswith('|'):
            table_buf.append(lines[i]); i+=1
        add_table(doc,table_buf); table_buf=[]; continue
    if not ln.strip() or ln.strip()=='---': i+=1; continue
    if ln.startswith('# '):
        # Skip duplicate main title because title page exists.
        if not ln.startswith('# NovaMart Guardian'): 
            doc.add_heading(ln[2:].strip(),0)
    elif ln.startswith('## '): doc.add_heading(ln[3:].strip(),1)
    elif ln.startswith('### '):
        if ln[4:].strip()=='Counts verified': doc.add_page_break()
        doc.add_heading(ln[4:].strip(),2)
    elif ln.startswith('#### '): doc.add_heading(ln[5:].strip(),3)
    elif ln.startswith('- '):
        p=doc.add_paragraph(style='List Bullet'); add_inline(p, ln[2:].strip())
    elif re.match(r'^\d+\.\s+',ln):
        p=doc.add_paragraph(); add_inline(p, ln)
    elif ln.startswith('> '):
        p=doc.add_paragraph(); r=p.add_run(ln[2:]); r.italic=True; r.font.color.rgb=RGBColor(100,40,40)
    elif ln.startswith('![Architecture]'):
        if png and png.exists(): doc.add_picture(str(png),width=Inches(6.7))
    else:
        p=doc.add_paragraph()
        add_inline(p, ln)
    i+=1

# Privacy / metadata cleanup
core=doc.core_properties
core.author='NovaMart Guardian Team'; core.title='NovaMart Guardian - Mantra Yudha Project Report'; core.subject='Agentic AI customer support project report'; core.comments=''
doc.save(out)
print(out)
